from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\Users\kuromi\Documents\ChatGPT\sxb\上行宝（SXB）项目解决方案.docx"
BLUE, DARK, LIGHT, GOLD = "2E74B5", "1F4D78", "E8EEF5", "7A5A00"

def set_font(run, size=10.5, bold=False, color=None):
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.font.size = Pt(size)
    run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr(); el = OxmlElement('w:shd'); el.set(qn('w:fill'), fill); pr.append(el)

def cell_margin(cell):
    pr = cell._tc.get_or_add_tcPr(); mar = OxmlElement('w:tcMar')
    for side in ('top','start','bottom','end'):
        el = OxmlElement('w:'+side); el.set(qn('w:w'), '100'); el.set(qn('w:type'),'dxa'); mar.append(el)
    pr.append(mar)

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers)); table.alignment = WD_TABLE_ALIGNMENT.LEFT; table.autofit = False
    if widths is None: widths = [9360 // len(headers)] * len(headers)
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths): col.set(qn('w:w'), str(width))
    for cell, text, width in zip(table.rows[0].cells, headers, widths):
        shade(cell, LIGHT); cell_margin(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        run = cell.paragraphs[0].add_run(text); set_font(run, 9.5, True, DARK)
        tcw = OxmlElement('w:tcW'); tcw.set(qn('w:w'), str(width)); tcw.set(qn('w:type'),'dxa'); cell._tc.get_or_add_tcPr().append(tcw)
    for values in rows:
        cells = table.add_row().cells
        for cell, text, width in zip(cells, values, widths):
            cell_margin(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); set_font(p.add_run(str(text)), 9.2)
            tcw = OxmlElement('w:tcW'); tcw.set(qn('w:w'), str(width)); tcw.set(qn('w:type'),'dxa'); cell._tc.get_or_add_tcPr().append(tcw)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(3); set_font(p.add_run(item))

def note(doc, title, text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(7)
    set_font(p.add_run(title+'：'), 10.5, True, GOLD); set_font(p.add_run(text))

doc = Document(); sec = doc.sections[0]
sec.top_margin, sec.bottom_margin = Inches(.8), Inches(.75)
sec.left_margin, sec.right_margin = Inches(.85), Inches(.85)
for name, size, color, before, after in [('Normal',10.5,None,0,6),('Heading 1',16,BLUE,16,8),('Heading 2',13,BLUE,12,6),('Heading 3',11.5,DARK,8,4)]:
    style = doc.styles[name]; style.font.name='Calibri'; style._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); style.font.size=Pt(size); style.font.bold=(name!='Normal'); style.paragraph_format.space_before=Pt(before); style.paragraph_format.space_after=Pt(after); style.paragraph_format.line_spacing=1.1
    if color: style.font.color.rgb=RGBColor.from_string(color)
header=sec.header.paragraphs[0]; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT; set_font(header.add_run('上行宝（SXB）项目解决方案 | 内部讨论稿'),8,False,'777777')
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_font(footer.add_run('上行宝（SXB） | 2026-08-08'),8,False,'777777')

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(38); p.paragraph_format.space_after=Pt(8); set_font(p.add_run('上行宝（SXB）项目解决方案'),25,True,DARK)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(26); set_font(p.add_run('面向成人考证的刷题驱动型全链路学习平台'),14,False,'505050')
add_table(doc,['文档属性','内容'],[['文档用途','项目架构与产品方案讨论'],['首个考试','初级社会工作师'],['首期形态','微信小程序 + H5 + 管理后台'],['方案状态','已确认需求与待讨论事项汇总']],[1800,7560])
note(doc,'阅读说明','已确认内容按当前需求执行；待讨论内容只做架构预留，不代表一期必须开发。')

doc.add_heading('一、项目概述',1)
doc.add_paragraph('上行宝是面向成人考证用户的刷题驱动型学习平台。平台围绕考试知识体系组织历年真题、课程、笔记、收藏、错题和挖孔背题，为用户提供从了解考试、建立知识框架、章节练习到考前巩固的完整学习路径。')
doc.add_heading('1.1 产品定位',2)
bullets(doc,['以刷题为核心入口，以知识点为组织单元，以学习进度和正确率为主要反馈。','每个考试拥有独立内容空间和独立购买权益。','一期不做复杂千人千面推荐，用户仅在当前考试上下文中学习。','首期以初级社会工作师验证内容模型和学习链路，后续通过配置扩展其他考试。'])
doc.add_heading('1.2 首期目标',2)
add_table(doc,['目标','范围'],[['用户端','微信小程序、H5；登录、考试切换、刷题、课程/知识点辅助、错题、收藏、笔记、学习计划'],['后台端','考试内容、知识点、题库、课程资源、用户、订单权益、角色权限'],['核心流程','购买/体验 → 进入当前考试 → 直接刷题；学习路径、课程、笔记、错题和报告作为辅助'],['后续扩展','广告、公告、PC、App、模拟考试、手动输入式挖孔、更多题型及更复杂 AI 能力']],[1800,7560])

doc.add_heading('二、已确认的业务边界',1)
add_table(doc,['主题','已确认规则'],[['考试与权益','用户可以同时购买多个考试；同一考试权益有效期内不可重复购买。'],['当前考试','当前考试界面只显示当前考试内容；首页和“我的”提供切换考试入口，消息中心除外。'],['考试名称','考试名称直接使用“初级社会工作师”等，不把年份写进考试名称。'],['服务期限','购买权益只能使用到当年考试结束；高级版、专业版到期后自动降级为基础版。'],['后台升级','工作人员可在后台人工升级用户权益，前端暂不提供升级流程。'],['题库年份','历年真题年份是题目筛选标签，不是知识结构层级，也不代表用户购买去年的考试。'],['版本等级','基础版、高级版、专业版先预留权益模型，具体差异待运营提供。'],['体验包','预留 1 元 12 小时基础版体验包。'],['题库来源','题库由后台工作人员导入，已有版权。'],['登录','支持手机号/短信验证码登录和微信授权登录；密码字段预留，首期前端不开放密码功能。']],[1800,7560])

doc.add_heading('三、产品信息架构',1)
doc.add_heading('3.1 用户端导航',2)
add_table(doc,['入口','主要内容'],[['首页','当前考试、倒计时、每日学习计划、今日题量、继续刷题'],['学习/刷题','科目、章、节、知识点、历年真题筛选、收藏题、错题本、挖孔背题'],['课程/知识内容','视频、音频、图文课程；知识点解析及关联入口'],['消息中心','系统通知、订单通知、考试相关消息；不受当前考试内容展示限制'],['我的','当前考试切换、考试权益、学习报告、学习笔记、收藏、错题本、订单与账户']],[1800,7560])
doc.add_heading('3.2 知识结构',2)
note(doc,'统一知识结构','考试类别 → 考试 → 科目 → 章 → 节 → 知识点。年份不放入知识树，作为题库标签和计划筛选条件。')
doc.add_paragraph('初级社会工作师首期配置两科：社会工作综合能力（初级）、社会工作实务（初级）。具体教材版本、大纲变化和考试时间由后台配置，并以当年官方信息为准。')

doc.add_heading('四、刷题与内容体系',1)
doc.add_heading('4.1 初级社会工作师一期题型',2)
doc.add_paragraph('一期重点支持单项选择题和多项选择题。旧需求文档中的 A1/A2、A3/A4、配伍题、主观题等属于未来跨考试扩展预留，不直接纳入初级社会工作师一期。')
add_table(doc,['题型','一期规则'],[['单项选择题','单题单选；提交后判定正误并展示总体解析。'],['多项选择题','支持多选答案；计分规则按初级社工考试要求维护。'],['挖孔背题','一期只做点击选项；不做手动输入。']],[2200,7160])
doc.add_heading('4.2 题目字段与关联',2)
bullets(doc,['题干、选项、标准答案、题型、总体解析。','关联一个主知识点，并可关联多个知识点。','关联所属考试、科目、章、节、知识点。','真题年份作为独立标签，用于筛选和学习计划。','支持收藏、错题本、作答记录、学习笔记入口。','题目由后台工作人员导入，支持草稿、上线、下线和导入错误反馈。'])
doc.add_heading('4.3 题目页面交互',2)
bullets(doc,['用户提交答案后立即显示正确/错误、标准答案和总体解析。','解析页提供知识点、课程、笔记入口。','用户可以收藏题目；答错题自动进入错题本。','用户可从章节、年份、收藏题、错题本等入口开始练习。','刷题完成状态按题目作答记录计算，并用于学习计划和报告。'])

doc.add_heading('五、刷题主导、路径辅助',1)
doc.add_paragraph('上行宝不是要求用户按固定步骤完成的课程产品。用户购买考试服务后，核心入口就是进入当前考试刷题；系统可以安排第 0 至第 3 轮的学习路径，向用户提供建议，但用户可以自由跳过、切换和重复练习。所有课程、知识点、笔记、收藏、错题、挖孔背题和学习报告，均服务于刷题提效。')
add_table(doc,['辅助阶段','辅助目标','对应能力'],[['第 0 轮：认识考试','帮助用户了解考什么','考试说明、知识地图、考试时间展示'],['第 1 轮：建立框架','帮助用户理解题目对应的知识点','视频/音频/图文课程、知识点内容、笔记'],['第 2 轮：强化纠错','帮助用户处理错题和薄弱点','错题本、收藏题、专项刷题、重点知识图谱'],['第 3 轮：冲刺巩固','帮助用户快速回顾重点','点击式挖孔背题、重点内容复习']], [2200,2900,4260])
doc.add_paragraph('学习路径是推荐和辅助，不是强制业务流程；用户任何时候都可以直接刷题。')

doc.add_heading('六、每日学习计划',1)
doc.add_paragraph('用户可以为当前考试设置每日学习计划。默认全选，用户可以调整题目来源。')
bullets(doc,['选择真题年份。','选择科目。','选择错题本内容。','选择收藏内容。','系统根据计划总题量和距离考试天数计算每日最低题量。','计划显示在首页，展示今日已完成、今日剩余、每日最低完成量和总体进度。'])
note(doc,'建议计算方式','每日最低题量 = 向上取整（计划总题量 ÷ 距离考试天数）。')
doc.add_paragraph('建议一期一个计划绑定一个考试，不把多个考试混在同一计划中；用户可在不同考试下分别创建计划。')

doc.add_heading('七、学习报告与 AI 分析',1)
doc.add_heading('7.1 已确认规则',2)
bullets(doc,['完成整个节、章或科目的全部题目后才生成对应报告。','用户退出刷题界面时弹出学习报告。','完成一个节时生成节报告；章和科目在自身全部完成后生成报告。','后续重新刷题导致掌握程度更新时，生成新的报告，不覆盖历史报告。','“我的”页面提供学习报告入口，可查看各科目、章、节的历史报告。','管理员从后台上传多张报告封面图，用户分享时选择封面。','模拟考试报告纳入待讨论项。'])
doc.add_heading('7.2 报告内容',2)
add_table(doc,['报告部分','内容'],[['学习统计','完成题数、正确题数、正确率、用时、完成日期'],['知识点分析','各知识点正确率、已掌握知识点、薄弱知识点'],['错题分析','错题数量、错题入口、重点复习知识点'],['学习建议','下一步课程、知识点和练习建议'],['分享内容','用户选择封面图后生成分享卡片；隐私字段需另行确认']],[2200,7160])
doc.add_heading('7.3 AI 调用方案',2)
doc.add_paragraph('系统先计算固定统计数据，再将结构化统计结果提交 DeepSeek 生成分析文字。报告保存原始统计数据、模型版本、提示词版本和 AI 返回内容，保证历史报告可追溯。AI 服务异常时仍展示固定统计数据，并显示分析暂不可用。')

doc.add_heading('八、文章、广告与公告（P2）',1)
doc.add_heading('8.1 广告',2)
bullets(doc,['广告整体列为 P2，不进入一期核心范围。','P2 预留后台上传图片/视频、考试隔离、H5/小程序投放、广告位和跳转配置。','未来 App 作为新增投放端。','一期不做曝光量和点击量统计。'])
doc.add_heading('8.2 文章与公告',2)
bullets(doc,['考试公告整体列为 P2，不进入一期核心范围。','文章类型预留政策、考试须知、公告、学习资料。','文章支持富文本、图片、音频、视频和附件。','公告必须关联考试；当前考试只显示当前考试公告。','文章和公告均支持置顶；展示形式预留普通列表、弹窗、公告条。'])

doc.add_heading('九、权益、订单与支付',1)
bullets(doc,['每个考试独立购买权益；有效期内同一考试不可重复购买。','产品等级基础版、高级版、专业版先做成可配置权益，不提前写死差异。','权益到考试结束日自动处理；高级版和专业版自动降级为基础版。','后台工作人员可以人工升级用户权益，记录操作人、时间、原因和有效期。','H5 支持支付宝手机网站支付；小程序支持微信支付。','订单必须经过支付回调验签后发放权益，并支持幂等、退款和异常状态。'])

doc.add_heading('十、后台管理方案',1)
add_table(doc,['模块','一期职责'],[['考试与知识结构','考试、科目、章、节、知识点的配置和排序'],['题库管理','固定题型题目录入、批量导入、解析、年份标签、知识点挂载、上下线'],['课程资源','视频、音频、图文课程及知识点关联'],['笔记与知识内容','官方知识点内容、课程关联、用户笔记入口配置'],['用户与权益','账户、考试权益、到期状态、后台人工升级'],['订单财务','订单查询、支付宝/微信支付状态、退款记录'],['角色权限','最高意志、小编、财务等预置角色；自定义角色和功能点授权'],['操作审计','后台新增、修改、上下线、权益调整和敏感操作日志']],[2200,7160])
doc.add_paragraph('旧需求中的模拟考试、推广大使、积分、广告统计、用户包、人群定向、复杂 AI 监控等内容暂不纳入一期，但后台权限和数据模型可以预留扩展位置。')

doc.add_heading('十一、技术架构建议',1)
doc.add_paragraph('当前团队主要由产品经理和运营组成，建议采用“单体应用 + 模块化设计”，减少部署和维护复杂度，避免一期拆分微服务。')
add_table(doc,['层次','建议方案'],[['用户端','uni-app：微信小程序、H5 共用主要页面与业务代码'],['后台端','Vue 3 + Element Plus，面向小编、财务和最高意志'],['后端','TypeScript + NestJS，统一 API、鉴权、业务模块和支付回调'],['数据库','MySQL，保存用户、考试、题目、知识点、作答、订单、权益和报告'],['缓存与任务','Redis，用于短信验证码、缓存、报告异步任务、支付重试'],['媒体资源','对象存储 + CDN，保存视频、音频、图片和报告封面'],['部署','Docker 部署到独立服务器，区分开发、测试、生产环境'],['AI 服务','DeepSeek API，后端统一调用，记录模型和提示词版本']],[2200,7160])
doc.add_heading('11.1 后端模块边界',2)
bullets(doc,['身份与用户模块','考试与知识结构模块','题库与作答模块','课程与媒体模块','笔记、收藏、错题模块','学习计划与进度模块','学习报告与 AI 模块','订单与权益模块','文章、公告与广告模块（P2）','后台角色权限与审计模块'])
doc.add_heading('11.2 关键架构原则',2)
bullets(doc,['考试内容和用户权益分离：同一考试内容可以被不同等级权益控制。','考试年份只作为题目标签，不进入知识树。','题型一期固定，但题目数据结构保留未来扩展能力。','用户端所有学习接口必须带当前考试上下文。','所有后台内容支持草稿、发布、下线和操作记录。','AI 失败不影响基础学习和统计功能。'])

doc.add_heading('十二、一期范围建议',1)
add_table(doc,['纳入一期','暂不纳入一期'],[['手机号/短信登录、微信授权登录','模拟考试'],['考试切换和当前考试隔离','广告与考试公告（P2）'],['知识结构和知识点内容','手动输入式挖孔背题'],['单选题、多选题和总体解析','自定义题型编辑器'],['视频、音频、图文课程','复杂掌握度算法'],['刷题、收藏、错题本、笔记','广告曝光/点击统计'],['点击式挖孔背题','推广大使、积分、用户包、广告定向'],['每日学习计划','PC、App'],['科目/章/节学习报告与 DeepSeek 分析','模拟考试相关报告']],[4680,4680])

doc.add_heading('十三、待讨论事项',1)
add_table(doc,['事项','当前状态','需要讨论的决策'],[['版本权益差异','待运营提供','基础版、高级版、专业版具体功能和价格'],['DeepSeek 报告生成','待确认','每次重新生成、同一轮只生成一次，还是允许人工重新生成'],['AI 降级策略','待确认','AI 不可用时是否只展示固定统计数据'],['报告分享隐私','待确认','是否展示昵称、头像、正确率、掌握等级'],['模拟考试','明确后续讨论','题型、组卷、计分、报告是否纳入学习数据'],['考试结束日','需配置','后台配置具体日期，还是接入官方日历'],['公告展示优先级','待确认','弹窗与公告条是否支持排序、频次和已读状态'],['计划完成定义','待确认','是否要求题目全部作答一次，还是达到完成比例']],[2200,2200,4960])

doc.add_heading('十四、架构结论',1)
doc.add_paragraph('上行宝一期应优先建设“考试内容可配置、题目可追溯、学习过程可记录、用户权益可控制”的基础平台。初级社会工作师只是首个内容实例，不应形成独立代码分支。考试、知识点、题目、课程、权益、广告、文章和报告都应通过考试 ID 进行隔离；历年真题年份作为题目标签参与筛选和计划，不参与考试购买。')
note(doc,'方案定位','这是面向产品和研发讨论的架构方案，不等同于最终 PRD、UI 设计稿或开发排期。')
doc.save(OUT)
print(OUT)
